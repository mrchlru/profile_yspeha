#!/usr/bin/env python3
"""Генерирует faqTestReferenceData.ts из xlsx «ТАБЛИЦА ТЕСТОВ справочно.xlsx»."""

from __future__ import annotations

import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = Path.home() / "Downloads" / "ТАБЛИЦА ТЕСТОВ справочно.xlsx"
OUT_PATH = ROOT / "src" / "lib" / "support" / "faqTestReferenceData.ts"

DASHBOARD_DETAIL_SHEET: dict[str, str] = {
    "cfit": "IQ Доп",
    "gerchikov": "Герчиков",
    "goal": "Целеустремленность",
    "paperwork": "Стиль работы с док",
    "rowe": "Направленность",
    "mono_multitask": "Моно",
    "adaptability": "Адаптивность",
    "risk": "Риск",
    "stress": "Стрессоустойчивость",
    "locus": "Локус",
    "thomas": "Томас",
    "kos": "КОС",
    "burnout": "ПИ",
    "burnout_pi": "ПИ",
}


def _cell_str(value: object | None) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
    return re.sub(r"\n{3,}", "\n\n", text)


def _parse_test_sheet(rows: list[tuple]) -> list[dict[str, str]]:
    parsed: list[dict[str, str]] = []
    current_name = ""
    for row in rows[1:]:
        name = _cell_str(row[0] if len(row) > 0 else None)
        prop = _cell_str(row[1] if len(row) > 1 else None)
        values = _cell_str(row[2] if len(row) > 2 else None)
        if name:
            current_name = name
        elif not current_name:
            continue
        if not prop and not values:
            continue
        parsed.append(
            {
                "testName": current_name,
                "propertyDescription": prop,
                "possibleValues": values,
            }
        )
    return parsed


def _parse_detail_sheet(rows: list[tuple]) -> list[str]:
    lines: list[str] = []
    for row in rows:
        parts = [_cell_str(cell) for cell in row if _cell_str(cell)]
        if not parts:
            continue
        if len(parts) == 1:
            lines.append(parts[0])
        else:
            lines.append(" — ".join(parts))
    return lines


def _ts_string(value: str) -> str:
    return json.dumps(value, ensure_ascii=False)


def build_module(xlsx_path: Path) -> str:
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    test_rows = _parse_test_sheet(list(wb["ТЕСТ"].iter_rows(values_only=True)))

    detail_sheets: dict[str, list[str]] = {}
    for sheet_name in wb.sheetnames:
        if sheet_name == "ТЕСТ":
            continue
        detail_sheets[sheet_name] = _parse_detail_sheet(
            list(wb[sheet_name].iter_rows(values_only=True))
        )
    wb.close()

    lines: list[str] = [
        "/** Автогенерация: web/scripts/buildFaqTestReferenceModule.py */",
        "",
        "export type FaqTestReferenceTableRow = {",
        "  testName: string;",
        "  propertyDescription: string;",
        "  possibleValues: string;",
        "};",
        "",
        "export type FaqTestReferenceDetailSheet = {",
        "  sheetTitle: string;",
        "  lines: ReadonlyArray<string>;",
        "};",
        "",
        "export const FAQ_TEST_REFERENCE_TABLE: ReadonlyArray<FaqTestReferenceTableRow> = [",
    ]
    for row in test_rows:
        lines.append(
            "  {"
            f" testName: {_ts_string(row['testName'])},"
            f" propertyDescription: {_ts_string(row['propertyDescription'])},"
            f" possibleValues: {_ts_string(row['possibleValues'])},"
            " },"
        )
    lines.append("];")
    lines.append("")
    lines.append(
        "export const FAQ_TEST_REFERENCE_DETAILS: ReadonlyArray<FaqTestReferenceDetailSheet> = ["
    )
    for sheet_name, sheet_lines in detail_sheets.items():
        lines.append(f"  {{ sheetTitle: {_ts_string(sheet_name)}, lines: [")
        for line in sheet_lines:
            lines.append(f"    {_ts_string(line)},")
        lines.append("  ]},")
    lines.append("];")
    lines.append("")
    lines.append(
        "export const FAQ_DASHBOARD_METHODOLOGY_DETAIL_SHEET: Readonly<"
        "Partial<Record<string, string>>"
        "> = {"
    )
    for info_id, sheet in DASHBOARD_DETAIL_SHEET.items():
        lines.append(f"  {_ts_string(info_id)}: {_ts_string(sheet)},")
    lines.append("};")
    lines.append("")
    return "\n".join(lines) + "\n"


def main() -> None:
    xlsx = DEFAULT_XLSX
    if not xlsx.is_file():
        raise SystemExit(f"Файл не найден: {xlsx}")
    OUT_PATH.write_text(build_module(xlsx), encoding="utf-8")
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
